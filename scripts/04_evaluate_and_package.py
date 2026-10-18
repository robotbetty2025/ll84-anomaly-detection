"""Step 4: per-type residual scaling, proxy-label benchmark, stability, SHAP, audit workbook, app artifacts."""
import json, sys, time
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import joblib, numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt, seaborn as sns
from scipy.stats import spearmanr
from sklearn.base import clone
from src import config as C
from src.baseline import PercentileBaseline
from src.evaluation import benchmark, build_pool, pairwise_stability, proxy_labels, topk_jaccard
from src.features import TARGET
from src.scoring import fit_type_scaling, scale_residuals

t0 = time.time(); sns.set_theme(style="whitegrid")
allr = pd.read_parquet(C.PROCESSED / "ll84_all_classified.parquet")
ready = allr[allr.status == "model_ready"]
train = ready[ready.year.isin(C.TRAIN_YEARS)].reset_index(drop=True)
oof = pd.read_parquet(C.PROCESSED / "train_oof.parquet")
bundle = joblib.load(C.MODELS / "lightgbm_residual.joblib")
pipe, feats, sigma_global = bundle["pipeline"], bundle["features"], bundle["sigma"]
pool = build_pool(allr, train.prop_type.unique())
unsup = pd.read_parquet(C.PROCESSED / "pool_scores_unsupervised.parquet")
assert (unsup.property_id.values == pool.property_id.values).all() and (unsup.year.values == pool.year.values).all()
pool["if_score"], pool["ocsvm_score"] = unsup.if_score.values, unsup.ocsvm_score.values
res = {}

# ---------- 1. LightGBM scores: global sigma vs per-type robust scaling
pool["pred_log"] = pipe.predict(pool[feats]); pool["pred_eui"] = np.exp(pool.pred_log)
pool["resid"] = pool[TARGET] - pool.pred_log
tab, glob = fit_type_scaling(oof.prop_type, oof[TARGET] - oof.oof_log_pred, k=50)
pool["score_global"] = pool.resid / sigma_global
pool["score_pertype"] = scale_residuals(pool.prop_type, pool.resid, tab, glob)
base = PercentileBaseline(0.90).fit(train)
pool["baseline_ratio"] = base.score(pool)
pool["baseline_flag"] = pool.baseline_ratio > 1
# two-sided variant of the rule (flags unusually LOW EUI as well) so the comparison with a two-sided residual score is fair
p10 = train.groupby("prop_type").site_eui.quantile(0.10)
pool["baseline2_ratio"] = np.maximum(pool.site_eui / pool.prop_type.map(base.thresholds_), pool.prop_type.map(p10) / pool.site_eui)
rd = pool[pool.status == "model_ready"]            # comparable to the week-4 hold-out set
big = rd.prop_type.value_counts().head(10).index
def flag_rates(score):
    thr = rd[score.name].quantile(0.90)
    return (rd.assign(f=rd[score.name] >= thr).groupby("prop_type").f.mean())
fr = pd.DataFrame({"global_sigma": flag_rates(rd.score_global), "per_type_sigma": flag_rates(rd.score_pertype),
                   "p90_rule": rd.groupby("prop_type").baseline_flag.mean(), "n": rd.prop_type.value_counts()}).loc[big]
res["flag_rate_by_type"] = fr.round(3).to_dict("index")
res["flag_rate_spread_sd"] = {"global_sigma": float(fr.global_sigma.std()), "per_type_sigma": float(fr.per_type_sigma.std())}
res["type_scaling_table_top"] = tab.loc[big].round(3).to_dict("index")
fr[["p90_rule", "global_sigma", "per_type_sigma"]].plot.barh(figsize=(8, 5), color=["#9aa5b1", "#d9822b", "#2a7f62"])
plt.xlabel("share flagged (2024 hold-out, top 10% overall)"); plt.title("Per-type scaling evens out flag rates")
plt.tight_layout(); plt.savefig(C.FIGS / "15_flag_rate_per_type_scaling.png", dpi=150); plt.close()
print("scoring done", round(time.time() - t0), flush=True)

# ---------- 2. proxy-label benchmark
labels = proxy_labels(pool)
res["proxy_label_counts"] = {k: int(v.sum()) for k, v in labels.items()} | {"pool": len(pool)}
scores = {"P90 rule (EUI / type P90)": pool.baseline_ratio.values,
          "P10/P90 rule (two-sided)": pool.baseline2_ratio.values,
          "EUI / national median": (pool.site_eui / pool.nat_median_site_eui).values,
          "LightGBM |resid| (global sigma)": pool.score_global.abs().values,
          "LightGBM |resid| (per-type sigma)": pool.score_pertype.abs().values,
          "Isolation Forest": pool.if_score.values, "One-Class SVM": pool.ocsvm_score.values}
bm = {}
for lab in ["any_proxy", "any_excl_plausibility", "plausibility_fail", "yoy_shift_gt3x", "extreme_es_score"]:
    bm[lab] = benchmark(labels[lab].values, scores).round(4)
    bm[lab].to_csv(C.REPORTS / f"benchmark_{lab}.csv", index=False)
res["benchmark"] = {k: v.to_dict("records") for k, v in bm.items()}
fig, axes = plt.subplots(1, 2, figsize=(12, 4.2))
for ax, lab, ttl in zip(axes, ["any_proxy", "any_excl_plausibility"], ["All proxy labels", "Excluding plausibility failures"]):
    b = bm[lab].set_index("method")[["pr_auc", "precision@10%"]]
    b.plot.barh(ax=ax, color=["#2a7f62", "#5b7db1"]); ax.axvline(bm[lab].base_rate.iloc[0], color="r", ls="--", lw=1)
    ax.set_title(f"{ttl}\n(base rate {bm[lab].base_rate.iloc[0]:.1%}, dashed line)", fontsize=10); ax.set_ylabel(""); ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.1), ncol=2, frameon=False)
plt.tight_layout(); plt.savefig(C.FIGS / "16_proxy_benchmark.png", dpi=150); plt.close()
print("benchmark done", round(time.time() - t0), flush=True)

# ---------- 3. inter-model agreement (model-ready 2024 only)
rd = pool[pool.status == "model_ready"]
m = pd.DataFrame({"P10/P90 rule": rd.baseline2_ratio, "LGBM per-type": rd.score_pertype.abs(), "Isolation Forest": rd.if_score, "OC-SVM": rd.ocsvm_score})
sp = m.corr(method="spearman").round(3)  # all four scores are two-sided
jc = pd.DataFrame({a: {b: topk_jaccard(m[a], m[b]) for b in m} for a in m}).round(3)
res["agreement_spearman"], res["agreement_top10_jaccard"] = sp.to_dict(), jc.to_dict()
fig, axes = plt.subplots(1, 2, figsize=(11, 4))
sns.heatmap(sp, annot=True, fmt=".2f", cmap="Blues", vmin=0, vmax=1, ax=axes[0]); axes[0].set_title("Spearman rank correlation")
sns.heatmap(jc, annot=True, fmt=".2f", cmap="Greens", vmin=0, vmax=1, ax=axes[1]); axes[1].set_title("Top-10% overlap (Jaccard)")
plt.tight_layout(); plt.savefig(C.FIGS / "17_model_agreement.png", dpi=150); plt.close()

# ---------- 4. seed stability: LightGBM and Isolation Forest
lg = []
for s in (11, 22, 33, 44, 55):
    p = clone(pipe).set_params(model__random_state=s).fit(train[feats], train[TARGET], model__categorical_feature=[0, 1])
    r = pool[TARGET] - p.predict(pool[feats]); lg.append(np.abs(scale_residuals(pool.prop_type, r, tab, glob)))
res["stability_lightgbm"] = pairwise_stability(lg)
res["stability_iforest"] = json.loads((C.REPORTS / "week5_unsupervised_log.json").read_text())["iforest_stability_final"]
print("stability done", round(time.time() - t0), flush=True)

# ---------- 5. SHAP (LightGBM TreeSHAP values via booster.predict(pred_contrib=True))
import shap
Xt_all = pipe[:-1].transform(pool[feats])
booster = pipe.named_steps["model"].booster_
rs = np.random.RandomState(C.RANDOM_STATE); idx = rs.choice(len(pool), 3000, replace=False)
contrib = booster.predict(Xt_all.iloc[idx], pred_contrib=True)
sv = contrib[:, :-1]
imp = pd.Series(np.abs(sv).mean(0), index=Xt_all.columns).sort_values(ascending=False)
res["shap_global_mean_abs"] = imp.round(4).to_dict()
plt.figure(); shap.summary_plot(sv, Xt_all.iloc[idx], show=False, max_display=12); plt.title("SHAP summary (LightGBM, 3,000 sampled 2024 buildings)", fontsize=10)
plt.tight_layout(); plt.savefig(C.FIGS / "18_shap_summary.png", dpi=150, bbox_inches="tight"); plt.close()
imp.head(12)[::-1].plot.barh(figsize=(7, 4.5), color="#2a7f62"); plt.xlabel("mean |SHAP| (effect on ln EUI)"); plt.title("Global SHAP feature importance")
plt.tight_layout(); plt.savefig(C.FIGS / "19_shap_global_importance.png", dpi=150); plt.close()
print("shap done", round(time.time() - t0), flush=True)

# ---------- 6. blind audit workbook (50 buildings: 25 top-ranked, 25 mid-ranked)
rd = pool[pool.status == "model_ready"].copy()
rd["abs_score"] = rd.score_pertype.abs(); rd["pct"] = rd.abs_score.rank(pct=True)
top = rd[rd.pct >= 0.98].sample(25, random_state=C.RANDOM_STATE).assign(group="top")
mid = rd[(rd.pct >= 0.45) & (rd.pct <= 0.55)].sample(25, random_state=C.RANDOM_STATE).assign(group="mid")
aud = pd.concat([top, mid]).sample(frac=1, random_state=7).reset_index(drop=True)
prev = allr[allr.year == 2023].drop_duplicates("property_id").set_index("property_id").site_eui
aud["audit_id"] = np.arange(1, 51)
aud["eui_2023"] = aud.property_id.map(prev)
cols = {"audit_id": "ID", "property_id": "Property ID", "property_name": "Name", "address": "Address", "borough": "Borough",
        "prop_type": "Property type", "year_built": "Year built", "gfa": "Floor area (ft2)", "eui_2023": "Site EUI 2023",
        "site_eui": "Site EUI 2024", "pred_eui": "Model expected EUI", "share_elec": "Electric share", "share_gas": "Gas share",
        "share_steam": "Steam share", "share_oil": "Oil share", "occupancy": "Occupancy %", "default_values": "Default values flag",
        "estimated_energy": "Estimated energy flag"}
sheet = aud[list(cols)].rename(columns=cols).round({"Floor area (ft2)": 0, "Site EUI 2023": 1, "Site EUI 2024": 1, "Model expected EUI": 1,
                                                    "Electric share": 2, "Gas share": 2, "Steam share": 2, "Oil share": 2})
key = aud[["audit_id", "group", "score_pertype", "score_global", "if_score", "baseline_ratio"]].round(3)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
wb = Workbook(); ws = wb.active; ws.title = "Audit"
head = list(sheet.columns) + ["Your classification", "Notes"]
ws.append(head)
for r in sheet.itertuples(index=False): ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r] + [None, None])
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2A7F62"); c.alignment = Alignment(wrap_text=True, vertical="center")
for i, h in enumerate(head, 1): ws.column_dimensions[get_column_letter(i)].width = 38 if h in ("Name", "Address", "Notes") else (24 if h == "Your classification" else 14)
ws.freeze_panes = "C2"; ws.row_dimensions[1].height = 32
dv = DataValidation(type="list", formula1='"Plausible Anomaly,Data Error,Normal"', allow_blank=True); ws.add_data_validation(dv)
dv.add(f"{get_column_letter(len(head)-1)}2:{get_column_letter(len(head)-1)}51")
ws2 = wb.create_sheet("Instructions", 0)
for line in ["50-building domain audit (blind)", "",
             "Rows are shuffled and the model's rank is hidden on purpose, so judge each building on its own merits.",
             "For each building choose one classification in the 'Your classification' column:",
             "  Plausible Anomaly = EUI is unusual and you can think of a real operational reason (inefficiency, odd use, etc.)",
             "  Data Error = the numbers look like a reporting mistake (unit error, wrong area, default values, impossible EUI)",
             "  Normal = nothing unusual for this property type, size and age",
             "Useful context to check: EUI change 2023 to 2024, fuel shares, default/estimated flags, the property name and address.",
             "When finished, save this file; the 'Key' sheet (do not open until done) shows which rows were top vs mid ranked."]:
    ws2.append([line])
ws2["A1"].font = Font(bold=True, size=14); ws2.column_dimensions["A"].width = 120
ws3 = wb.create_sheet("Key (open after auditing)"); ws3.append(list(key.columns))
for r in key.itertuples(index=False): ws3.append(list(r))
wb.save(C.REPORTS / "expert_audit_50_buildings.xlsx")
res["audit_groups"] = key.group.value_counts().to_dict()
print("audit workbook done", round(time.time() - t0), flush=True)

# ---------- 7. app artifacts
A = C.ROOT / "app" / "data"; A.mkdir(parents=True, exist_ok=True)
rd = pool[pool.status == "model_ready"].copy()
thr_lgbm = float(rd.score_pertype.quantile(0.90)); thr_if = float(rd.if_score.quantile(0.90))
rd["lgbm_flag"] = rd.score_pertype >= thr_lgbm; rd["if_flag"] = rd.if_score >= thr_if
keep = ["property_id", "property_name", "address", "borough", "prop_type", "year_built", "gfa", "site_eui", "pred_eui", "resid",
        "score_global", "score_pertype", "baseline_ratio", "baseline_flag", "lgbm_flag", "if_score", "if_flag", "es_score"] + [f for f in feats if f not in ("prop_type", "borough", "log_site_eui")] + ["log_site_eui"]
keep = list(dict.fromkeys(keep))
rd[keep].reset_index(drop=True).to_parquet(A / "scored_2024.parquet", index=False)
peers = train.groupby("prop_type").site_eui.describe(percentiles=[.1, .25, .5, .75, .9])[["count", "10%", "25%", "50%", "75%", "90%"]]
ifb = joblib.load(C.MODELS / "isolation_forest.joblib")
joblib.dump({"lgbm": pipe, "features": feats, "type_scaling": tab, "global_scaling": glob, "sigma_global": sigma_global,
             "baseline_thresholds": base.thresholds_, "iforest": ifb["pipeline"], "iforest_features": ifb["features"],
             "thr_lgbm": thr_lgbm, "thr_if": thr_if, "if_ref": np.sort(rd.if_score.values),
             "nat_median_by_type": train.groupby("prop_type").nat_median_site_eui.median(), "peer_stats": peers,
             "prop_types": sorted(train.prop_type.unique()), "boroughs": sorted(train.borough.dropna().unique())}, A / "model_bundle.joblib", compress=3)
res["thresholds"] = {"lgbm_pertype_top10": thr_lgbm, "iforest_top10": thr_if}
w4 = json.loads((C.REPORTS / "week4_metrics.json").read_text())
res["regression_eval_2024"] = {"property_type_median": w4["ref_type_median"], "linear_regression": w4["ref_linear_regression"], "lightgbm": w4["lgbm_test_2024"]}
res["lgbm_best_params"], res["lgbm_cv_rmse_log"] = w4["lgbm_best_params"], w4["lgbm_cv_rmse_log"]
(C.REPORTS / "final_metrics.json").write_text(json.dumps(res, indent=2, default=str)); (A / "results.json").write_text(json.dumps(res, default=str))
import shutil
for f in ["03_eui_by_type_box", "04_eui_vs_gfa", "12_lgbm_actual_vs_expected", "13_flag_rate_baseline_vs_lgbm", "15_flag_rate_per_type_scaling", "16_proxy_benchmark", "17_model_agreement", "18_shap_summary", "19_shap_global_importance"]:
    (A.parent / "figures").mkdir(exist_ok=True); shutil.copy(C.FIGS / f"{f}.png", A.parent / "figures" / f"{f}.png")
print(json.dumps({k: res[k] for k in ["flag_rate_spread_sd", "proxy_label_counts", "stability_lightgbm", "stability_iforest", "audit_groups"]}, indent=1))
print("total", round(time.time() - t0), "s")
