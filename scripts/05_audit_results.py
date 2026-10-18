"""Step 5: score the completed 50-building audit (precision of top-ranked vs mid-ranked buildings).

Run after filling in the 'Your classification' column of reports/expert_audit_50_buildings.xlsx.
Exits with a message (and writes nothing) if the workbook is incomplete.
"""
import json, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import numpy as np, pandas as pd
import matplotlib; matplotlib.use("Agg")
import matplotlib.pyplot as plt, seaborn as sns
from openpyxl import load_workbook
from scipy.stats import fisher_exact
from src import config as C

CLASSES = ["Plausible Anomaly", "Data Error", "Normal"]
wb = load_workbook(C.REPORTS / "expert_audit_50_buildings.xlsx", data_only=True)
ws = wb["Audit"]; hdr = [c.value for c in ws[1]]
rows = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
aud = pd.DataFrame(rows)[["ID", "Your classification", "Notes"]]
key = pd.DataFrame(list(wb["Key (open after auditing)"].iter_rows(values_only=True))[1:],
                   columns=[c.value for c in wb["Key (open after auditing)"][1]])
done = aud["Your classification"].isin(CLASSES)
if not done.all():
    print(f"Audit incomplete: {int(done.sum())}/50 classified. Fill in every row (dropdown values only) and re-run."); sys.exit(2)
d = aud.merge(key, left_on="ID", right_on="audit_id")
ct = pd.crosstab(d.group, d["Your classification"]).reindex(index=["top", "mid"], columns=CLASSES, fill_value=0)
d["non_normal"] = d["Your classification"] != "Normal"
top, mid = d[d.group == "top"], d[d.group == "mid"]
t22 = np.array([[top.non_normal.sum(), (~top.non_normal).sum()], [mid.non_normal.sum(), (~mid.non_normal).sum()]])
odds, p = fisher_exact(t22)
res = {"counts": ct.to_dict("index"), "n_top": len(top), "n_mid": len(mid),
       "precision_top_plausible_anomaly": float((top["Your classification"] == "Plausible Anomaly").mean()),
       "precision_top_any_non_normal": float(top.non_normal.mean()),
       "data_error_share_top": float((top["Your classification"] == "Data Error").mean()),
       "mid_plausible_anomaly_share": float((mid["Your classification"] == "Plausible Anomaly").mean()),
       "mid_any_non_normal": float(mid.non_normal.mean()),
       "fisher_exact_p_top_vs_mid_non_normal": float(p), "odds_ratio": float(odds)}
(C.REPORTS / "audit_results.json").write_text(json.dumps(res, indent=2))
ct.to_csv(C.REPORTS / "audit_confusion_counts.csv")
fig, axes = plt.subplots(1, 2, figsize=(10, 3.2))
sns.heatmap(ct, annot=True, fmt="d", cmap="Greens", cbar=False, ax=axes[0]); axes[0].set_ylabel("Model rank group"); axes[0].set_xlabel("Auditor classification"); axes[0].set_title("Audit results (50 buildings)")
b = pd.DataFrame({"Flagged (top 2%)": [top.non_normal.sum(), (~top.non_normal).sum()], "Not flagged (mid)": [mid.non_normal.sum(), (~mid.non_normal).sum()]}, index=["Anomalous (plausible or data error)", "Normal"])
sns.heatmap(b, annot=True, fmt="d", cmap="Blues", cbar=False, ax=axes[1]); axes[1].set_title("Precision confusion matrix")
plt.tight_layout(); plt.savefig(C.FIGS / "20_audit_confusion.png", dpi=150); plt.close()
print(json.dumps(res, indent=2))
